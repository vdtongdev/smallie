from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from app.features.documents import images
from app.features.documents.images import images_in


def save_image(path: Path, size: tuple[int, int], mode: str = "RGB", color=(200, 30, 30)) -> Path:
    Image.new(mode, size, color).save(path)
    return path


def opened(data: bytes) -> Image.Image:
    return Image.open(BytesIO(data))


def make_pdf(path: Path, pages: list[list[Path]]) -> Path:
    """One PDF page per entry, each with some text and the given images."""
    pdf = canvas.Canvas(str(path), pagesize=A4)
    for page_images in pages:
        pdf.drawString(72, 800, "Some text")
        for i, image in enumerate(page_images):
            pdf.drawImage(ImageReader(str(image)), 72, 600 - i * 160, 200, 150)
        pdf.showPage()
    pdf.save()
    return path


def test_image_file_becomes_one_rgb_jpeg(tmp_path):
    path = save_image(tmp_path / "scan.jpg", (800, 600), "CMYK", (0, 128, 255, 0))
    [(page, data)] = images_in(path)
    assert page is None
    assert opened(data).format == "JPEG" and opened(data).mode == "RGB"


def test_images_with_a_short_side_under_100px_are_skipped(tmp_path):
    assert images_in(save_image(tmp_path / "icon.png", (50, 50))) == []
    assert images_in(save_image(tmp_path / "strip.png", (99, 600))) == []
    assert len(images_in(save_image(tmp_path / "ok.png", (100, 600)))) == 1


def test_large_images_are_shrunk_to_1600px(tmp_path):
    [(_, data)] = images_in(save_image(tmp_path / "photo.png", (4000, 3000)))
    assert opened(data).size == (1600, 1200)


def test_pdf_images_keep_their_page(tmp_path):
    picture = save_image(tmp_path / "p.png", (400, 300))
    pdf = make_pdf(tmp_path / "r.pdf", [[], [picture]])
    assert [page for page, _ in images_in(pdf)] == [2]


def test_pdf_without_images_has_none(tmp_path):
    assert images_in(make_pdf(tmp_path / "plain.pdf", [[]])) == []


def test_docx_images_have_no_page_and_skip_the_thumbnail(tmp_path):
    doc = DocxDocument()
    doc.add_paragraph("Hello")
    doc.add_picture(str(save_image(tmp_path / "a.png", (400, 300), color=(255, 0, 0))))
    doc.add_picture(str(save_image(tmp_path / "b.png", (400, 300), color=(0, 0, 255))))
    doc.save(tmp_path / "n.docx")
    assert [page for page, _ in images_in(tmp_path / "n.docx")] == [None, None]


def test_at_most_max_images_per_file(tmp_path, monkeypatch):
    monkeypatch.setattr(images, "MAX_IMAGES_PER_FILE", 2)
    pictures = [save_image(tmp_path / f"{i}.png", (400, 300), color=(i * 80, 0, 0)) for i in range(3)]
    assert len(images_in(make_pdf(tmp_path / "many.pdf", [pictures]))) == 2


def test_unreadable_image_is_skipped_not_fatal(tmp_path):
    path = tmp_path / "broken.png"
    path.write_bytes(b"this is not a png")
    assert images_in(path) == []
