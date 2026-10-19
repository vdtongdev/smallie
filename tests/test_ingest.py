from pathlib import Path

import pytest
from docx import Document as DocxDocument
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.features.documents import ingest
from app.features.documents.images import ImageReading
from app.features.documents.ingest import chunk_ids, load_and_split

VI_TEXT = "Chính sách đổi trả: khách hàng được đổi sản phẩm trong vòng 30 ngày kể từ ngày mua."
READING = ImageReading(description="A bar chart of sales.", text="Q1 10, Q2 12", doc_type="chart",
                       summary="Quarterly sales.", entities=[])


@pytest.fixture
def english_pdf(tmp_path: Path) -> Path:
    path = tmp_path / "handbook.pdf"
    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.drawString(72, 800, "Page one: the office opens at 8 am.")
    pdf.showPage()
    pdf.drawString(72, 800, "Page two: the refund window is 30 days.")
    pdf.showPage()
    pdf.save()
    return path


@pytest.fixture
def vietnamese_docx(tmp_path: Path) -> Path:
    path = tmp_path / "chinh-sach.docx"
    doc = DocxDocument()
    doc.add_heading("Quy định cửa hàng", level=1)
    doc.add_paragraph(VI_TEXT)
    doc.save(path)
    return path


def test_pdf_chunks_keep_filename_and_one_based_pages(english_pdf: Path):
    chunks = load_and_split(english_pdf, "handbook.pdf")
    assert [c.metadata["page"] for c in chunks] == [1, 2]
    assert all(c.metadata["filename"] == "handbook.pdf" for c in chunks)
    assert "refund window" in chunks[1].page_content


def test_docx_keeps_vietnamese_diacritics(vietnamese_docx: Path):
    chunks = load_and_split(vietnamese_docx, "chính-sách.docx")
    text = " ".join(c.page_content for c in chunks)
    assert VI_TEXT in text
    assert chunks[0].metadata == {"filename": "chính-sách.docx"}


def test_long_text_is_split_into_bounded_chunks(tmp_path: Path):
    path = tmp_path / "long.docx"
    doc = DocxDocument()
    for _ in range(40):
        doc.add_paragraph(VI_TEXT)
    doc.save(path)
    chunks = load_and_split(path, "long.docx")
    assert len(chunks) > 1
    assert max(len(c.page_content) for c in chunks) <= 800


def test_unsupported_type_rejected(tmp_path: Path):
    path = tmp_path / "notes.txt"
    path.write_text("hello")
    with pytest.raises(ValueError):
        load_and_split(path, "notes.txt")


def test_chunk_ids_are_deterministic():
    assert chunk_ids("doc", 3) == ["doc:0", "doc:1", "doc:2"]


def test_image_chunks_come_after_text_chunks(english_pdf: Path):
    chunks = load_and_split(english_pdf, "handbook.pdf", [(2, READING)])
    assert [c.metadata.get("kind") for c in chunks] == [None, None, "image"]  # ids 0..n stay deterministic
    assert chunks[2].metadata["page"] == 2 and chunks[2].page_content.startswith("[Image 1, p. 2]")


def test_image_file_has_only_image_chunks_and_long_text_is_split(tmp_path: Path):
    long_reading = READING.model_copy(update={"text": VI_TEXT * 40})
    chunks = load_and_split(tmp_path / "scan.png", "scan.png", [(None, long_reading)])
    assert len(chunks) > 1
    assert all(c.metadata["kind"] == "image" and len(c.page_content) <= 800 for c in chunks)
    assert load_and_split(tmp_path / "scan.png", "scan.png") == []  # no reading, no chunks


def test_ingest_saves_image_count_and_returns_readings(tmp_path: Path, monkeypatch):
    saved, stored = {}, []
    monkeypatch.setattr(ingest, "read_file_images", lambda path, filename, callbacks: [(None, READING)])
    monkeypatch.setattr(ingest, "_set_status", lambda document_id, status, **values: saved.update(status=status, **values))
    monkeypatch.setattr(ingest, "get_vector_store", lambda: type("Store", (), {
        "add_documents": lambda self, docs, ids: stored.extend(ids)})())
    chunks, readings = ingest.ingest_document("doc", "u1", tmp_path / "chart.png", "chart.png", ["usage"])
    assert readings == [READING]
    assert saved == {"status": "ready", "chunk_count": 1, "image_count": 1, "error": None}
    assert stored == ["doc:0"] and chunks[0].metadata["user_id"] == "u1"


def test_image_that_could_not_be_read_fails_the_file(tmp_path: Path, monkeypatch):
    saved = {}
    monkeypatch.setattr(ingest, "read_file_images", lambda path, filename, callbacks: [])  # every call failed
    monkeypatch.setattr(ingest, "_set_status", lambda document_id, status, **values: saved.update(status=status, **values))
    assert ingest.ingest_document("doc", "u1", tmp_path / "photo.jpg", "photo.jpg") == ([], [])
    assert saved["status"] == "failed" and saved["error"] == "No text or readable images found."
